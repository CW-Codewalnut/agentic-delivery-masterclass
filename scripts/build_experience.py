#!/usr/bin/env python3
"""Build the offline masterclass HTML and native deck from one scene source."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCENES = ROOT / "presentation" / "scenes.json"
ROLE_PATHS = {
    "capture-refine": ("roles/01-capture-refine.md", "skills/capture-refine/SKILL.md"),
    "design": ("roles/02-design.md", "skills/design/SKILL.md"),
    "planner": ("roles/03-planner.md", "skills/planner/SKILL.md"),
    "builder": ("roles/04-builder.md", "skills/builder/SKILL.md"),
    "tester": ("roles/05-tester.md", "skills/tester/SKILL.md"),
    "reviewer": ("roles/06-reviewer.md", "skills/reviewer/SKILL.md"),
    "curator": ("roles/07-curator.md", "skills/curator/SKILL.md"),
}
EXTRA_PATHS = ["domain/order-cancellation.md", "docs/worked-example.md"]
DOWNLOADS = {
    "Masterclass-Experience.pdf": "presentation/Masterclass-Experience.pdf",
    "Masterclass-Experience.pptx": "presentation/Masterclass-Experience.pptx",
    "Presenter-Script.md": "presenter/SCRIPT.md",
    "Presenter-Cue-Sheet.md": "presenter/CUE-SHEET.md",
    "Presenter-Timing.md": "presenter/TIMING.md",
}


def read_sources(source_root: Path) -> tuple[dict[str, dict], list[str]]:
    files: dict[str, dict] = {}
    missing: list[str] = []
    supporting: list[str] = []
    for pattern in ("worked-example/**/*", "tests/*.py", "docs/*.md", "README.md"):
        supporting.extend(str(path.relative_to(source_root)) for path in source_root.glob(pattern) if path.is_file())
    paths = list(dict.fromkeys([p for pair in ROLE_PATHS.values() for p in pair] + EXTRA_PATHS + sorted(supporting)))
    for path in paths:
        candidate = source_root / path
        if candidate.is_file():
            content = candidate.read_text(encoding="utf-8")
            files[path] = {
                "path": path,
                "content": content,
                "sha256": hashlib.sha256(content.encode()).hexdigest(),
                "bytes": len(content.encode()),
            }
        else:
            missing.append(path)
    return files, missing


def html_document(data: dict, embedded: dict[str, dict], missing: list[str]) -> str:
    payload = json.dumps({"experience": data, "files": embedded, "missing": missing}, ensure_ascii=False).replace("</", "<\\/")
    return f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="theme-color" content="#071016"><title>{html.escape(data['meta']['title'])} — Masterclass</title>
<style>
:root{{--ink:#071016;--panel:#0e1c25;--panel2:#152a36;--paper:#f2f0e9;--muted:#a9b7bf;--cyan:#58d2df;--green:#84d6a2;--line:#294653;--amber:#edca72;--space:clamp(16px,2vw,32px)}}
*{{box-sizing:border-box}}html,body{{height:100%}}body{{margin:0;background:var(--ink);color:var(--paper);font:16px/1.45 Avenir Next,Avenir,Inter,system-ui,sans-serif;overflow:hidden}}
button{{font:inherit}}button:focus-visible{{outline:3px solid var(--paper);outline-offset:3px}}.skip{{position:fixed;left:12px;top:-80px;z-index:30;background:var(--paper);color:var(--ink);padding:12px 16px}}.skip:focus{{top:12px}}
.shell{{height:100%;display:grid;grid-template-rows:auto 1fr auto}}header{{display:flex;align-items:center;gap:16px;padding:14px var(--space);border-bottom:1px solid var(--line)}}.brand{{font-size:11px;font-weight:700;letter-spacing:.16em;color:var(--cyan)}}.status{{margin-left:auto;font-size:12px;color:var(--muted)}}.status strong{{color:var(--green)}}
main{{min-height:0;display:grid;grid-template-columns:220px 1fr}}.rail{{padding:20px 14px;border-right:1px solid var(--line);overflow:auto}}.rail-label{{display:block;color:var(--muted);font-size:10px;letter-spacing:.15em;margin:0 8px 10px}}.role{{width:100%;display:grid;grid-template-columns:24px 1fr;text-align:left;gap:9px;padding:11px 8px;border:0;border-radius:4px;background:transparent;color:var(--muted);cursor:pointer}}.role:hover{{background:var(--panel)}}.role[aria-current=true]{{background:var(--panel2);color:var(--paper)}}.role b{{font-size:13px}}.role small{{display:block;font-size:11px;color:var(--muted)}}.role-num{{font:17px/1 Georgia,serif;color:var(--cyan)}}
.stage{{position:relative;min-width:0;overflow:hidden}}.scene{{position:absolute;inset:0;padding:clamp(28px,4vw,64px);display:none;grid-template-rows:auto auto 1fr auto;gap:18px}}.scene.active{{display:grid}}.kicker{{color:var(--cyan);font-size:11px;font-weight:700;letter-spacing:.16em}}h1{{font:700 clamp(34px,4.2vw,68px)/.98 Georgia,Times,serif;letter-spacing:-.025em;margin:0;max-width:1050px}}.lede{{font-size:clamp(16px,1.45vw,22px);color:var(--muted);max-width:900px;margin:0}}.body{{display:flex;align-items:center;min-height:0}}.quote{{font:clamp(30px,4vw,60px)/1.08 Georgia,serif;color:var(--cyan);max-width:900px}}.cards{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;width:100%}}.cards.two{{grid-template-columns:repeat(2,minmax(0,1fr));gap:18px}}.card{{position:relative;background:var(--panel);padding:20px;min-height:180px;border-radius:5px;overflow:hidden}}.card:before{{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:var(--cyan)}}.card.final:before,.card.supported:before{{background:var(--green)}}.card.unassessed:before{{background:var(--amber)}}.card.stale:before{{background:var(--muted)}}.card-label{{font-size:10px;font-weight:700;letter-spacing:.14em;color:var(--cyan)}}.card h2{{font:700 clamp(19px,2vw,29px)/1.08 Avenir Next,system-ui,sans-serif;margin:26px 0 12px}}.card p{{margin:0;color:var(--muted);font-size:14px}}.questions{{width:100%;display:grid;grid-template-columns:1fr 1fr;gap:10px 30px}}.question{{display:grid;grid-template-columns:84px 1fr;align-items:start;padding:18px 0;border-top:1px solid var(--line)}}.question b{{font-size:12px;color:var(--cyan);letter-spacing:.12em}}.question span{{font-size:19px}}.takeaway{{background:var(--panel2);padding:18px 22px;border-radius:5px;font:clamp(18px,1.8vw,27px)/1.25 Georgia,serif}}
footer{{display:grid;grid-template-columns:1fr auto;gap:20px;align-items:center;padding:12px var(--space);border-top:1px solid var(--line);background:var(--ink)}}.acts{{display:grid;grid-template-columns:repeat(6,1fr);gap:6px}}.act{{border:0;border-top:3px solid var(--line);background:transparent;color:var(--muted);text-align:left;padding:7px 4px 0;font-size:10px;line-height:1.25;cursor:pointer}}.act.active{{border-color:var(--cyan);color:var(--paper)}}.controls{{display:flex;align-items:center;gap:8px}}.control{{border:1px solid var(--line);background:transparent;color:var(--paper);border-radius:4px;padding:9px 12px;cursor:pointer}}.control.primary{{background:var(--cyan);border-color:var(--cyan);color:var(--ink);font-weight:700}}.control:disabled{{opacity:.35;cursor:not-allowed}}.count{{min-width:56px;text-align:center;font-variant-numeric:tabular-nums;color:var(--muted);font-size:12px}}
.filebar{{position:absolute;right:18px;top:18px;display:flex;gap:8px;z-index:4}}.file-trigger,.download-link{{border:1px solid var(--line);background:rgba(7,16,22,.88);color:var(--paper);padding:8px 11px;border-radius:4px;font-size:11px;cursor:pointer;text-decoration:none}}.drawer{{position:fixed;inset:0 0 0 auto;width:min(720px,92vw);z-index:20;background:#0a151c;border-left:1px solid var(--line);box-shadow:-24px 0 80px #0008;transform:translateX(102%);transition:transform .25s ease;display:grid;grid-template-rows:auto auto 1fr}}.drawer.open{{transform:none}}.drawer-head{{display:flex;align-items:center;padding:18px 22px;border-bottom:1px solid var(--line)}}.drawer-head h2{{font:28px Georgia,serif;margin:0}}.drawer-close{{margin-left:auto}}.downloads{{display:flex;flex-wrap:wrap;gap:7px;padding:12px 22px 0}}.downloads a{{color:var(--cyan);font-size:12px}}.file-tabs{{display:flex;gap:8px;padding:12px 22px;overflow:auto;border-bottom:1px solid var(--line)}}.file-tab{{white-space:nowrap;border:1px solid var(--line);background:transparent;color:var(--muted);padding:8px 10px;border-radius:4px;cursor:pointer}}.file-tab[aria-selected=true]{{background:var(--cyan);color:var(--ink);border-color:var(--cyan)}}.file-view{{margin:0;padding:24px;overflow:auto;white-space:pre-wrap;font:13px/1.65 ui-monospace,SFMono-Regular,Menlo,monospace;color:#d7e1e5}}.file-meta{{color:var(--muted);font-size:11px;padding:0 22px}}.screen-note{{position:fixed;left:-9999px}}
@media(max-width:850px){{main{{grid-template-columns:1fr}}.rail{{display:none}}.scene{{padding:26px 20px}}.cards,.questions{{grid-template-columns:1fr 1fr}}.card{{min-height:135px}}.acts{{display:none}}.status{{max-width:45%;text-align:right}}}}
@media(max-width:560px){{header{{padding:10px 14px}}.brand{{max-width:52%}}.status{{font-size:11px}}.stage{{padding-top:54px}}.filebar{{top:5px;right:16px}}.scene.active{{display:block}}.scene{{top:54px;padding:16px 16px 96px;overflow-y:auto}}.scene .kicker{{display:block;margin-bottom:14px}}.scene .lede{{margin-top:12px}}.scene .body{{display:block;min-height:0;margin:20px 0}}h1{{font-size:34px;line-height:1.03}}.cards,.cards.two,.questions{{grid-template-columns:1fr}}.question{{grid-template-columns:72px 1fr}}.card{{min-height:0}}.takeaway{{font-size:18px;margin-top:16px}}footer{{padding:9px 12px}}.control,.file-trigger{{min-height:44px;padding:9px 12px}}}}
@media(prefers-reduced-motion:reduce){{*{{transition:none!important}}}}
@media print{{body{{overflow:visible;background:white;color:#071016}}header,.rail,footer,.filebar,.drawer{{display:none!important}}main{{display:block}}.stage{{overflow:visible}}.scene{{position:relative;opacity:1;transform:none;pointer-events:auto;height:7.5in;break-after:page;padding:.5in}}.scene:not(.active){{display:grid}}.lede,.card p{{color:#43535c}}.kicker,.quote,.question b,.card-label{{color:#24515c!important}}.card{{background:#edf1f2}}.takeaway{{background:#e4ecef}}}}
</style></head>
<body><a class="skip" href="#stage">Skip to presentation</a><div class="shell">
<header><span class="brand">{html.escape(data['meta']['eyebrow'])}</span><span class="status"><strong>SYNTHETIC EXAMPLE · ILLUSTRATIVE POLICY</strong> · prepared offline walkthrough</span></header>
<main><nav class="rail" aria-label="Role journey"><span class="rail-label">SEVEN ACCOUNTABLE ROLES</span><div id="roles"></div></nav><section class="stage" id="stage" tabindex="-1" aria-label="Presentation scenes"><div class="filebar"><a class="download-link" href="downloads/Masterclass-Experience.pdf" download>PDF</a><a class="download-link" href="downloads/Masterclass-Experience.pptx" download>PPTX</a><button class="file-trigger" id="openFiles" data-control="open-files">Source & presenter files</button></div><div id="scenes"></div></section></main>
<footer><nav class="acts" id="acts" aria-label="Six acts"></nav><div class="controls"><button class="control" id="reset" data-control="reset" title="Restart presentation (Home)">Reset</button><button class="control" id="back" data-control="back" title="Back (←)">← Back</button><span class="count" id="count"></span><button class="control primary" id="next" data-control="next" title="Next (→ or Space)">Next →</button></div></footer></div>
<aside class="drawer" id="drawer" role="dialog" aria-modal="true" aria-hidden="true" inert aria-labelledby="drawerTitle"><div class="drawer-head"><h2 id="drawerTitle">Repository source</h2><button class="control drawer-close" id="closeFiles" data-control="close-files">Close</button></div><div><nav class="downloads" aria-label="Downloads"><a href="downloads/Masterclass-Experience.pdf" download>Deck PDF</a><a href="downloads/Masterclass-Experience.pptx" download>Deck PPTX</a><a href="downloads/Presenter-Script.md" download>Presenter script</a><a href="downloads/Presenter-Cue-Sheet.md" download>Cue sheet</a><a href="downloads/Presenter-Timing.md" download>Timing</a></nav><div class="file-tabs" id="fileTabs" role="tablist" aria-label="Embedded source files"></div><div class="file-meta" id="fileMeta"></div></div><pre class="file-view" id="fileView" tabindex="0"></pre></aside><div class="screen-note" id="announce" aria-live="polite"></div>
<script id="masterclass-data" type="application/json">{payload}</script>
<script>
(()=>{{'use strict';const DATA=JSON.parse(document.getElementById('masterclass-data').textContent),X=DATA.experience,S=X.scenes,R=X.roles,F=DATA.files;let i=0,lastFocus=null;
const $=id=>document.getElementById(id),esc=s=>{{const d=document.createElement('div');d.textContent=s;return d.innerHTML}},scenes=$('scenes'),roles=$('roles'),acts=$('acts');
function cards(scene){{const cls=scene.kind==='split'?'cards two':'cards';return `<div class="${{cls}}">${{(scene.items||[]).map((x,n)=>{{const semantic=['SUPPORTED','UNASSESSED','STALE'].includes(x[0])?x[0].toLowerCase():'';return `<article class="card ${{n===scene.items.length-1?'final':''}} ${{semantic}}"><span class="card-label">${{esc(x[0])}}</span><h2>${{esc(x[1]||'')}}</h2>${{x[2]?`<p>${{esc(x[2])}}</p>`:''}}</article>`}}).join('')}}</div>`}}
function body(s){{if(s.kind==='hero')return `<div class="quote">${{esc(s.quote)}}</div>`;if(s.kind==='questions')return `<div class="questions">${{s.items.map(x=>`<div class="question"><b>${{esc(x[0])}}</b><span>${{esc(x[1])}}</span></div>`).join('')}}</div>`;return cards(s)}}
S.forEach((s,n)=>{{const el=document.createElement('article');el.className='scene';el.id='scene-'+s.id;el.dataset.index=n;el.setAttribute('aria-hidden','true');el.innerHTML=`<span class="kicker">${{esc(s.kicker)}}</span><div><h1>${{esc(s.title)}}</h1><p class="lede">${{esc(s.lede)}}</p></div><div class="body">${{body(s)}}</div><div class="takeaway">${{esc(s.takeaway)}}</div>`;scenes.appendChild(el)}});
R.forEach((r,n)=>{{const b=document.createElement('button');b.className='role';b.dataset.role=r.id;b.dataset.control='role-'+r.id;b.innerHTML=`<span class="role-num">${{String(n+1).padStart(2,'0')}}</span><span><b>${{esc(r.name)}}</b><small>${{esc(r.verb)}}</small></span>`;b.onclick=()=>go(Math.max(0,S.findIndex(s=>s.role===r.id)));roles.appendChild(b)}});
X.acts.forEach((a,n)=>{{const b=document.createElement('button');b.className='act';b.dataset.act=n;b.dataset.control='act-'+(n+1);b.textContent=(n+1)+' / '+a;b.onclick=()=>go(S.findIndex(s=>s.act===n));acts.appendChild(b)}});
function go(n,push=true){{i=Math.max(0,Math.min(S.length-1,n));document.querySelectorAll('.scene').forEach((e,k)=>{{e.classList.toggle('active',k===i);e.setAttribute('aria-hidden',k===i?'false':'true')}});document.querySelectorAll('.role').forEach(e=>e.setAttribute('aria-current',e.dataset.role===S[i].role?'true':'false'));document.querySelectorAll('.act').forEach(e=>e.classList.toggle('active',+e.dataset.act===S[i].act));$('back').disabled=i===0;$('next').disabled=i===S.length-1;$('count').textContent=`${{String(i+1).padStart(2,'0')}} / ${{String(S.length).padStart(2,'0')}}`;$('announce').textContent=`Scene ${{i+1}}: ${{S[i].title}}`;if(push)history.replaceState(null,'','#'+S[i].id)}}
$('next').onclick=()=>go(i+1);$('back').onclick=()=>go(i-1);$('reset').onclick=()=>go(0);
function candidateFiles(){{const role=R.find(r=>r.id===S[i].role),preferred=[];if(role)preferred.push(role.roleFile,role.skillFile);X.artifacts.forEach(a=>preferred.push(a.path));return [...new Set([...preferred,...Object.keys(F)])].filter(p=>F[p])}}
function showFile(path){{document.querySelectorAll('.file-tab').forEach(e=>e.setAttribute('aria-selected',e.dataset.path===path?'true':'false'));$('fileView').textContent=F[path].content;$('fileMeta').textContent=`${{path}} · ${{F[path].bytes}} bytes · SHA-256 ${{F[path].sha256}}`}}
function openDrawer(){{lastFocus=document.activeElement;const paths=candidateFiles().length?candidateFiles():Object.keys(F);$('fileTabs').innerHTML='';paths.forEach((p,n)=>{{const b=document.createElement('button');b.className='file-tab';b.dataset.path=p;b.setAttribute('role','tab');b.textContent=p;b.onclick=()=>showFile(p);$('fileTabs').appendChild(b)}});$('drawer').classList.add('open');$('drawer').removeAttribute('inert');$('drawer').setAttribute('aria-hidden','false');if(paths[0])showFile(paths[0]);else{{$('fileView').textContent='Source files were not present at build time. Rebuild with --source-root after the repository content is available.';$('fileMeta').textContent=DATA.missing.length+' expected files unavailable';}}$('closeFiles').focus()}}
function closeDrawer(){{$('drawer').classList.remove('open');$('drawer').setAttribute('aria-hidden','true');$('drawer').setAttribute('inert','');if(lastFocus)lastFocus.focus()}}$('openFiles').onclick=openDrawer;$('closeFiles').onclick=closeDrawer;
document.addEventListener('keydown',e=>{{if($('drawer').classList.contains('open')){{if(e.key==='Escape'){{closeDrawer();return}}if(e.key==='Tab'){{const focusable=[...$('drawer').querySelectorAll('button,a[href],[tabindex]:not([tabindex="-1"])')].filter(x=>!x.disabled);const first=focusable[0],last=focusable[focusable.length-1];if(e.shiftKey&&document.activeElement===first){{e.preventDefault();last.focus()}}else if(!e.shiftKey&&document.activeElement===last){{e.preventDefault();first.focus()}}}}return}}if(['ArrowRight','PageDown'].includes(e.key)||(e.key===' '&&!/BUTTON|A/.test(e.target.tagName))){{e.preventDefault();go(i+1)}}else if(['ArrowLeft','PageUp'].includes(e.key)){{e.preventDefault();go(i-1)}}else if(e.key==='Home'){{e.preventDefault();go(0)}}else if(e.key.toLowerCase()==='f')openDrawer()}});
window.addEventListener('hashchange',()=>{{const n=S.findIndex(s=>s.id===location.hash.slice(1));if(n>=0&&n!==i)go(n,false)}});
const hash=location.hash.slice(1),start=S.findIndex(s=>s.id===hash);go(start>=0?start:0,false);
}})();
</script></body></html>'''


def build_pptx(data: dict, output: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise SystemExit("python-pptx is required for deck generation; install requirements.txt") from exc
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    bg, panel, paper, muted, cyan, green = [RGBColor(*x) for x in [(7,16,22),(14,28,37),(242,240,233),(169,183,191),(88,210,223),(132,214,162)]]
    def box(slide,x,y,w,h,color,radius=False):
        from pptx.enum.shapes import MSO_SHAPE
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid();sh.fill.fore_color.rgb=color;sh.line.fill.background();return sh
    def txt(slide,x,y,w,h,text,size=18,color=paper,bold=False,font="Avenir Next",align=None):
        sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
        p=tf.paragraphs[0];p.text=text;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
        if align is not None:p.alignment=align
        return sh
    for n,s in enumerate(data["scenes"]):
        slide=prs.slides.add_slide(blank);slide.background.fill.solid();slide.background.fill.fore_color.rgb=bg
        txt(slide,.5,.18,7.2,.2,data["meta"]["eyebrow"],8,cyan,True)
        txt(slide,8.0,.18,4.75,.2,"SYNTHETIC EXAMPLE · ILLUSTRATIVE POLICY",7,muted,True,align=PP_ALIGN.RIGHT)
        txt(slide,.5,.62,12,.3,s["kicker"],9,cyan,True)
        wraps = len(s["title"]) > 47
        subtitle_y = 2.03 if wraps else 1.82
        content_y = 2.75 if wraps else 2.55
        txt(slide,.5,1.02,12.2,1.02,s["title"],28 if wraps else 31,paper,True,"Georgia")
        txt(slide,.5,subtitle_y,11.7,.5,s["lede"],15,muted)
        items=s.get("items",[])
        if s["kind"]=="hero":
            txt(slide,.55,content_y+.2,11.8,1.25,s.get("quote",""),36,cyan,False,"Georgia")
        elif s["kind"]=="questions":
            for k,(label,body) in enumerate(items):
                x=.6+(k%2)*6.15;y=content_y+.1+(k//2)*1.05
                txt(slide,x,y,1,.3,label,10,cyan,True);txt(slide,x+1.05,y,4.8,.65,body,16,paper)
        else:
            cols=2 if s["kind"]=="split" else len(items); rows=(len(items)+cols-1)//cols
            gap=.18;w=(12.25-(cols-1)*gap)/cols;h=2.55/rows
            for k,item in enumerate(items):
                x=.5+(k%cols)*(w+gap);y=content_y+(k//cols)*(h+.12);box(slide,x,y,w,h,panel,True);box(slide,x,y,.045,h,green if k==len(items)-1 else cyan)
                txt(slide,x+.2,y+.18,w-.4,.2,item[0],8,cyan,True);txt(slide,x+.2,y+.58,w-.4,.55,item[1] if len(item)>1 else "",17,paper,True)
                if len(item)>2:txt(slide,x+.2,y+1.32,w-.4,.7,item[2],11,muted)
        box(slide,.5,5.55,12.25,.82,RGBColor(21,42,54),True);txt(slide,.72,5.75,11.8,.4,s["takeaway"],17,paper,False,"Georgia")
        for a in range(6):box(slide,.5+a*2.02,6.73,1.88,.025,cyan if a==s["act"] else RGBColor(41,70,83))
        txt(slide,12.15,6.85,.6,.25,f"{n+1:02d}",10,cyan,True,align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text=s.get("notes","")
    prs.core_properties.title=data["meta"]["title"]
    prs.core_properties.subject="Prepared walkthrough — seven-role agent-assisted SDLC"
    prs.core_properties.author="Ben Sheridan-Edwards"
    output.parent.mkdir(parents=True,exist_ok=True);prs.save(output)


def convert_pdf(pptx: Path, pdf: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: raise SystemExit("LibreOffice/soffice not found")
    pdf.parent.mkdir(parents=True,exist_ok=True)
    result=subprocess.run([soffice,"--headless","--convert-to","pdf","--outdir",str(pdf.parent),str(pptx)],capture_output=True,text=True,timeout=180)
    if result.returncode: raise SystemExit(result.stderr or result.stdout)
    generated=pdf.parent/(pptx.stem+".pdf")
    if generated != pdf: generated.replace(pdf)


def main() -> None:
    p=argparse.ArgumentParser();p.add_argument("--source-root",type=Path,default=ROOT);p.add_argument("--slice",action="store_true");p.add_argument("--skip-deck",action="store_true");a=p.parse_args()
    data=json.loads(SCENES.read_text())
    if a.slice:data={**data,"scenes":data["scenes"][:4]}
    files,missing=read_sources(a.source_root)
    contract={
        "source":"presentation/scenes.json",
        "controls":{
            "next":"button or ArrowRight/PageDown/Space",
            "back":"button or ArrowLeft/PageUp",
            "reset":"button or Home",
            "sourceDrawer":"button or F; Escape closes",
            "roleJump":"seven labelled role buttons",
            "actJump":"six exact Nattu act buttons"
        },
        "scenes":[{"index":n+1,"id":s["id"],"act":data["acts"][s["act"]],"role":s.get("role"),"title":s["title"],"kicker":s["kicker"],"takeaway":s["takeaway"],"presenterNotes":s.get("notes","")} for n,s in enumerate(data["scenes"])]
    }
    contract_path=ROOT/"presentation"/("slice-scene-control-contract.json" if a.slice else "scene-control-contract.json")
    contract_path.write_text(json.dumps(contract,ensure_ascii=False,indent=2)+"\n")
    web=ROOT/"web"/('capture-refine-slice.html' if a.slice else 'index.html');web.parent.mkdir(parents=True,exist_ok=True);web.write_text(html_document(data,files,missing),encoding="utf-8")
    outputs={"html":str(web.relative_to(ROOT)),"contract":str(contract_path.relative_to(ROOT)),"scenes":len(data["scenes"]),"embedded_files":len(files),"missing":missing}
    if not a.skip_deck:
        pptx=ROOT/"presentation"/("Capture-Refine-Slice.pptx" if a.slice else "Masterclass-Experience.pptx");pdf=pptx.with_suffix('.pdf');build_pptx(data,pptx);convert_pdf(pptx,pdf);outputs.update(pptx=str(pptx.relative_to(ROOT)),pdf=str(pdf.relative_to(ROOT)))
        if not a.slice:
            download_dir=ROOT/"web"/"downloads";download_dir.mkdir(parents=True,exist_ok=True)
            for name,source in DOWNLOADS.items(): shutil.copy2(ROOT/source,download_dir/name)
            outputs["downloads"]=[str((download_dir/name).relative_to(ROOT)) for name in DOWNLOADS]
    manifest=ROOT/"qa"/("slice-build.json" if a.slice else "build-manifest.json");manifest.parent.mkdir(parents=True,exist_ok=True);manifest.write_text(json.dumps(outputs,indent=2)+"\n")
    print(json.dumps(outputs))
if __name__=="__main__":main()
